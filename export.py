from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

def export_docx(title, sections, path):
    doc = Document()
    doc.add_heading(title, 0)

    for s in sections:
        doc.add_heading(s["heading"], level=1)
        doc.add_paragraph(s["text"])

    doc.save(path)


def export_pdf(title, sections, path):
    doc = SimpleDocTemplate(path)
    styles = getSampleStyleSheet()

    content = [Paragraph(title, styles["Title"])]

    for s in sections:
        content.append(Paragraph(s["heading"], styles["Heading2"]))
        content.append(Paragraph(s["text"], styles["BodyText"]))

    doc.build(content)
